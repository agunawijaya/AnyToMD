"""Flattened, geometry-aware view of a slide's shapes.

Shared by the PPTX layout code (reading order, titles, overlays), the
flowchart extractor and slide vision. Positions are in slide EMU, with group
transforms applied — python-pptx reports a grouped shape's offset in the
group's own child coordinate space.
"""
from __future__ import annotations

import re
from dataclasses import dataclass

from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

_A = "http://schemas.openxmlformats.org/drawingml/2006/main"

_PLACEHOLDER_ROLES = {
    PP_PLACEHOLDER.TITLE: "title",
    PP_PLACEHOLDER.CENTER_TITLE: "title",
    PP_PLACEHOLDER.VERTICAL_TITLE: "title",
    PP_PLACEHOLDER.SUBTITLE: "subtitle",
    PP_PLACEHOLDER.SLIDE_NUMBER: "number",
    PP_PLACEHOLDER.FOOTER: "footer",
    PP_PLACEHOLDER.DATE: "date",
}


@dataclass
class ShapeInfo:
    shape: object
    sid: str          # cNvPr id — what connectors refer to
    z: int            # drawing order; a later shape is drawn on top
    kind: str         # text | picture | table | chart | connector | shape | other
    left: float
    top: float
    right: float
    bottom: float
    text: str         # all text on one line: soft breaks and paragraphs as spaces
    prst: str         # preset geometry ("rect", "roundRect", "downArrow", …) or ""
    role: str         # placeholder role: title, subtitle, number, footer, date, body, or ""
    # connector ends in slide EMU (only for kind == "connector")
    begin: tuple[float, float] | None = None
    end: tuple[float, float] | None = None

    @property
    def width(self) -> float:
        return self.right - self.left

    @property
    def height(self) -> float:
        return self.bottom - self.top

    @property
    def area(self) -> float:
        return max(0.0, self.width) * max(0.0, self.height)

    @property
    def cx(self) -> float:
        return (self.left + self.right) / 2

    @property
    def cy(self) -> float:
        return (self.top + self.bottom) / 2

    def contains(self, x: float, y: float, tol: float = 0.0) -> bool:
        return (self.left - tol <= x <= self.right + tol
                and self.top - tol <= y <= self.bottom + tol)

    def overlap_share(self, other: "ShapeInfo") -> float:
        """Share of this shape's area that lies inside `other`."""
        w = min(self.right, other.right) - max(self.left, other.left)
        h = min(self.bottom, other.bottom) - max(self.top, other.top)
        if w <= 0 or h <= 0 or self.area <= 0:
            return 0.0
        return w * h / self.area


def clean_text(text: str) -> str:
    """Soft line breaks (`\\v`), paragraph breaks and runs of spaces → one space."""
    return re.sub(r"\s+", " ", (text or "").replace("\v", " ")).strip()


def collect_shapes(slide) -> list[ShapeInfo]:
    """Every shape on the slide, groups flattened, in drawing order."""
    out: list[ShapeInfo] = []
    _walk(slide.shapes, _identity, out)
    return out


def _identity(x: float, y: float) -> tuple[float, float]:
    return x, y


def _walk(shapes, transform, out: list[ShapeInfo]) -> None:
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            _walk(shape.shapes, _group_transform(shape, transform), out)
            continue
        info = _info(shape, transform, len(out))
        if info is not None:
            out.append(info)


def _group_transform(group, outer):
    """Child coordinates → slide coordinates for the shapes inside `group`."""
    xfrm = group._element.xpath("./p:grpSpPr/a:xfrm")
    if not xfrm:
        return outer
    x = xfrm[0]

    def _pair(tag: str, a: str, b: str) -> tuple[float, float] | None:
        node = x.find(f"{{{_A}}}{tag}")
        if node is None:
            return None
        return float(node.get(a, 0)), float(node.get(b, 0))

    off, ext = _pair("off", "x", "y"), _pair("ext", "cx", "cy")
    ch_off, ch_ext = _pair("chOff", "x", "y"), _pair("chExt", "cx", "cy")
    if not (off and ext and ch_off and ch_ext) or not ch_ext[0] or not ch_ext[1]:
        return outer
    sx, sy = ext[0] / ch_ext[0], ext[1] / ch_ext[1]

    def _t(px: float, py: float) -> tuple[float, float]:
        return outer(off[0] + (px - ch_off[0]) * sx, off[1] + (py - ch_off[1]) * sy)

    return _t


def _info(shape, transform, z: int) -> ShapeInfo | None:
    el = shape._element
    try:
        left, top = float(shape.left or 0), float(shape.top or 0)
        width, height = float(shape.width or 0), float(shape.height or 0)
    except Exception:
        return None
    l, t = transform(left, top)
    r, b = transform(left + width, top + height)
    sid = el.xpath("./*/p:cNvPr/@id")
    prst = el.xpath("./p:spPr/a:prstGeom/@prst")
    role = ""
    try:
        if shape.is_placeholder:
            role = _PLACEHOLDER_ROLES.get(shape.placeholder_format.type, "body")
    except Exception:
        role = ""

    begin = end = None
    if el.tag.endswith("}cxnSp") or shape.shape_type == MSO_SHAPE_TYPE.LINE:
        kind = "connector"
        try:
            begin = transform(float(shape.begin_x), float(shape.begin_y))
            end = transform(float(shape.end_x), float(shape.end_y))
        except Exception:
            begin, end = (l, t), (r, b)
    elif el.tag.endswith("}pic"):
        kind = "picture"
    elif getattr(shape, "has_table", False) and shape.has_table:
        kind = "table"
    elif getattr(shape, "has_chart", False) and shape.has_chart:
        kind = "chart"
    elif getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
        kind = "text"
    elif el.tag.endswith("}sp"):
        kind = "shape"
    else:
        kind = "other"

    text = ""
    if kind == "text":
        text = clean_text(shape.text_frame.text)
    return ShapeInfo(shape=shape, sid=str(sid[0]) if sid else "", z=z, kind=kind,
                     left=min(l, r), top=min(t, b), right=max(l, r), bottom=max(t, b),
                     text=text, prst=str(prst[0]) if prst else "", role=role,
                     begin=begin, end=end)
