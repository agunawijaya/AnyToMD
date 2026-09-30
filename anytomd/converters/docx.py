"""DOCX -> Markdown.

Headers, footers, and page-number fields live outside `<w:body>`, so iterating
the body naturally skips them. We walk the body in document order, preserving
paragraph vs. table interleaving, and describe embedded images via the AI.
"""
from __future__ import annotations

import re
from pathlib import Path
from typing import Callable, Iterator

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from ..ai_client import AIClient
from ..lang import detect_language

ProgressFn = Callable[[str], None]

_HEADING_RE = re.compile(r"^heading\s+(\d)$", re.IGNORECASE)


def convert_docx(path: Path, ai: AIClient, progress: ProgressFn | None = None) -> str:
    doc = Document(str(path))
    out: list[str] = [f"# {path.stem}", ""]
    language = detect_language(" ".join(p.text for p in doc.paragraphs))

    img_counter = 0
    para_counter = 0
    table_counter = 0

    for tag, elem in _iter_body(doc):
        if tag == "p":
            para_counter += 1
            para = Paragraph(elem, doc)
            heading_level = _heading_level(para)
            text = _paragraph_text(para)

            if text:
                if heading_level:
                    prefix = "#" * min(heading_level + 1, 6)  # +1 so h1 stays under file title
                    out.append(f"{prefix} {text}")
                elif _is_list_item(para):
                    out.append(f"- {text}")
                else:
                    out.append(text)
                out.append("")

            for rid in _drawing_rids(elem):
                img_counter += 1
                _emit(progress, f"Gambar {img_counter}: {rid}")
                blob, ctype = _load_image(doc, rid)
                if blob is None:
                    out.append(f"> *[Gagal ekstrak gambar {img_counter}]*")
                    out.append("")
                    continue
                try:
                    analysis = ai.analyze_image(
                        blob,
                        media_type=ctype or "image/png",
                        extra_context=("Gambar ini dari sebuah DOCX (inline pada dokumen). "
                                       f"Output language: {language}."),
                    )
                except Exception as e:
                    out.append(f"> *[Gagal analisa gambar {img_counter}: {e}]*")
                    out.append("")
                    continue
                block = analysis.to_markdown_block(alt_hint=f"docx_img{img_counter}")
                if block:  # silent: logo, icon, repeat
                    out.append(block)
                    out.append("")

        elif tag == "tbl":
            table_counter += 1
            _emit(progress, f"Tabel {table_counter}")
            table = Table(elem, doc)
            md = _table_to_md(table)
            if md:
                out.append(md)
                out.append("")

    return "\n".join(out).rstrip() + "\n"


# ---------- body walking ----------

def _iter_body(doc) -> Iterator[tuple[str, object]]:
    """Yield (local_tag, element) for each direct child of <w:body> in order."""
    body = doc.element.body
    for child in body.iterchildren():
        tag = child.tag.split("}", 1)[-1] if "}" in child.tag else child.tag
        if tag in {"p", "tbl"}:
            yield tag, child
        # sectPr and other structural elements: skip silently


def _heading_level(para: Paragraph) -> int:
    try:
        style_name = (para.style.name or "").lower()
    except Exception:
        return 0
    m = _HEADING_RE.match(style_name)
    if m:
        return int(m.group(1))
    if style_name == "title":
        return 1
    return 0


def _is_list_item(para: Paragraph) -> bool:
    """True if the paragraph is a list item (direct numPr OR list-style)."""
    pPr = para._p.find(qn("w:pPr"))
    if pPr is not None and pPr.find(qn("w:numPr")) is not None:
        return True
    try:
        style_name = (para.style.name or "").lower()
    except Exception:
        return False
    return "list" in style_name or style_name.startswith("bullet")


def _paragraph_text(para: Paragraph) -> str:
    text = para.text or ""
    # Collapse internal whitespace runs; python-docx already joins runs but may
    # keep multiple spaces from tabs etc.
    return re.sub(r"[ \t]+", " ", text).strip()


# ---------- image extraction ----------

def _drawing_rids(p_elem) -> list[str]:
    """Return rIds of blip-fill images inside this paragraph."""
    rids: list[str] = []
    for blip in p_elem.findall(".//" + qn("a:blip")):
        rid = blip.get(qn("r:embed"))
        if rid:
            rids.append(rid)
    return rids


def _load_image(doc, rid: str) -> tuple[bytes | None, str | None]:
    try:
        part = doc.part.related_parts[rid]
        return part.blob, getattr(part, "content_type", None)
    except (KeyError, AttributeError):
        return None, None


# ---------- tables ----------

def _table_to_md(table: Table) -> str:
    rows_text: list[list[str]] = []
    for row in table.rows:
        cells = [_cell_text(c) for c in row.cells]
        rows_text.append(cells)
    if not rows_text:
        return ""
    # Ensure uniform column count
    ncols = max(len(r) for r in rows_text)
    rows_text = [r + [""] * (ncols - len(r)) for r in rows_text]

    header, *body = rows_text
    lines = [
        "| " + " | ".join(_escape_md_cell(c) for c in header) + " |",
        "| " + " | ".join("---" for _ in header) + " |",
    ]
    for r in body:
        lines.append("| " + " | ".join(_escape_md_cell(c) for c in r) + " |")
    return "\n".join(lines)


def _cell_text(cell) -> str:
    # A cell can contain multiple paragraphs; join with a space so table rows
    # stay one line each (Markdown table cells can't wrap safely otherwise).
    parts = [p.text.strip() for p in cell.paragraphs if p.text.strip()]
    return " ".join(parts)


def _escape_md_cell(text: str) -> str:
    return text.replace("|", "\\|").replace("\n", " ").strip()


def _emit(progress: ProgressFn | None, msg: str) -> None:
    if progress:
        progress(msg)
