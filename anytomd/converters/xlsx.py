"""XLSX -> Markdown. All sheets combined into one MD, one per heading."""
from __future__ import annotations

import datetime as _dt
from pathlib import Path
from typing import Any, Callable

from openpyxl import load_workbook

from ..ai_client import AIClient  # not used for XLSX, but kept for uniform signature

ProgressFn = Callable[[str], None]

MAX_ROWS_PER_SHEET = 5000
MAX_COLS_PER_SHEET = 100


def convert_xlsx(path: Path, ai: AIClient | None = None,
                 progress: ProgressFn | None = None) -> str:
    wb = load_workbook(str(path), data_only=True)
    out: list[str] = [f"# {path.stem}", ""]

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        _emit(progress, f"Sheet: {sheet_name}")
        rows = _extract_rows(ws)
        out.append(f"## Sheet: {sheet_name}")
        out.append("")
        if not rows:
            out.append("*(sheet kosong)*")
            out.append("")
            continue

        truncated = False
        if len(rows) > MAX_ROWS_PER_SHEET:
            rows = rows[:MAX_ROWS_PER_SHEET]
            truncated = True
        if rows and len(rows[0]) > MAX_COLS_PER_SHEET:
            rows = [r[:MAX_COLS_PER_SHEET] for r in rows]
            truncated = True

        out.append(_rows_to_md(rows))
        out.append("")
        if truncated:
            out.append(f"*(dipotong pada {MAX_ROWS_PER_SHEET} baris × {MAX_COLS_PER_SHEET} kolom)*")
            out.append("")

    wb.close()
    return "\n".join(out).rstrip() + "\n"


# ---------- extraction ----------

def _extract_rows(ws) -> list[list[Any]]:
    all_rows = [list(row) for row in ws.iter_rows(values_only=True)]

    # Trim trailing empty rows
    while all_rows and all(_is_empty(v) for v in all_rows[-1]):
        all_rows.pop()
    if not all_rows:
        return []

    # Trim trailing empty columns
    last_col = 0
    for r in all_rows:
        for i in range(len(r) - 1, -1, -1):
            if not _is_empty(r[i]):
                last_col = max(last_col, i)
                break
    trimmed = [r[: last_col + 1] for r in all_rows]

    # Trim leading empty rows (whole rows only, keep structure otherwise)
    while trimmed and all(_is_empty(v) for v in trimmed[0]):
        trimmed.pop(0)

    return trimmed


def _is_empty(v: Any) -> bool:
    if v is None:
        return True
    if isinstance(v, str) and not v.strip():
        return True
    return False


# ---------- md rendering ----------

def _rows_to_md(rows: list[list[Any]]) -> str:
    if not rows:
        return ""
    ncols = max(len(r) for r in rows)
    rows = [r + [None] * (ncols - len(r)) for r in rows]

    header = rows[0]
    body = rows[1:]

    header_line = "| " + " | ".join(_fmt(c) for c in header) + " |"
    sep_line = "| " + " | ".join("---" for _ in header) + " |"
    lines = [header_line, sep_line]
    for r in body:
        lines.append("| " + " | ".join(_fmt(c) for c in r) + " |")
    return "\n".join(lines)


def _fmt(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        # Trim trailing zeros for integer-valued floats
        if v.is_integer():
            return str(int(v))
        return f"{v:g}"
    if isinstance(v, (_dt.datetime, _dt.date, _dt.time)):
        return v.isoformat()
    s = str(v)
    return s.replace("|", "\\|").replace("\n", " ").strip()


def _emit(progress: ProgressFn | None, msg: str) -> None:
    if progress:
        progress(msg)
