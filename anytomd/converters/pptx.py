"""PPTX → Markdown.

Two paths per slide, both built on the same exact pieces taken from the XML
(`pptx_layout`, `pptx_shape_mermaid`):

- **Deterministic** (always available; used for local models and as the
  fallback): `## <slide title>`, then the slide's blocks in reading order —
  paragraphs with their real bullets and numbering, native tables and charts
  as Markdown tables, flowcharts as Mermaid straight from the connectors,
  pictures analysed one by one (tables and pictures of text transcribed,
  logos and icons dropped).
- **Slide vision** (`pptx_slide_vision`, auto = cloud provider): the slide
  is rendered and a vision model writes its Markdown — headings, panels,
  columns — anchored on the XML text. Tables, pictures, charts and
  flowcharts enter the text layer as `[[PICTURE n]]`-style placeholders and
  are filled in afterwards from the exact pieces above, so their content
  never depends on the model re-reading pixels.
"""
from __future__ import annotations

import re
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable

from pptx import Presentation

from .. import pptx_render
from ..ai_client import PAGE_MIN_COVERAGE, AIClient, format_text_layer_entry, wrap_with_options
from ..lang import detect_language
from ..options import ConvertOptions
from ..slide_render import render_slide_shapes
from .pptx_layout import (Block, find_overlays, find_title, frame_markdown, is_page_number,
                          reading_order, table_markdown)
from .pptx_shape_mermaid import Flow, extract_flow, resolve_edges
from .pptx_shapes import ShapeInfo, collect_shapes

ProgressFn = Callable[[str], None]

# Alt text below this length is treated as auto-generated / non-informative
# ("Picture 3", "image1", empty). Above → trust the author.
_ALT_TEXT_MIN_LEN = 12
# Alt text Office writes itself, not the author: descriptions from its
# auto-alt-text feature, and file names.
_GENERATED_ALT_RE = re.compile(
    r"description automatically generated|deskripsi dibuat secara otomatis"
    r"|beschreibung automatisch generiert|^[\w\s().-]+\.(?:png|jpe?g|gif|bmp|svg|emf|wmf|tiff?|webp)$",
    re.IGNORECASE,
)
# Icons inserted from Office's icon library are SVG pictures named after the
# icon: "Bullseye outline", "Handshake with solid fill".
_OFFICE_ICON_ALT_RE = re.compile(r"^[\w\s'&,-]{1,40}\b(?:outline|with solid fill)$", re.IGNORECASE)
# How much slide body text we hand to the AI as extra_context (prevents
# extra_context from ballooning while still giving useful hints).
_CONTEXT_TEXT_MAX = 300
_SLIDE_VISION_WORKERS = 4
_TOKEN_KINDS = {"picture": "PICTURE", "table": "TABLE", "chart": "CHART", "diagram": "DIAGRAM"}
_TOKEN_LINE_RE = re.compile(r"^[ \t>*_`-]*\[\[(PICTURE|TABLE|CHART|DIAGRAM) (\d+)\]\][ \t*_`.]*$", re.M)
_TOKEN_RE = re.compile(r"\[\[(PICTURE|TABLE|CHART|DIAGRAM) (\d+)\]\]")
_CONNECTORS_HEADER = "CONNECTORS (exact, from the slide XML; tail → head):"


@dataclass
class _Slide:
    idx: int
    slide: object
    infos: list[ShapeInfo]
    title: ShapeInfo | None = None
    flow: Flow | None = None
    blocks: list[Block] = field(default_factory=list)
    shape_diagram: bool = False   # a shape diagram whose graph did not resolve
    notes: str = ""

    @property
    def title_text(self) -> str:
        return self.title.text if self.title else ""


def convert_pptx(path: Path, ai: AIClient, progress: ProgressFn | None = None,
                 options: ConvertOptions | None = None) -> str:
    prs = Presentation(str(path))
    slide_w, slide_h = prs.slide_width, prs.slide_height
    ai_images = wrap_with_options(ai, options) if options is not None else ai
    slides = [_prepare(i, s, slide_w, slide_h) for i, s in enumerate(prs.slides, start=1)]
    language = detect_language(" ".join(i.text for s in slides for i in s.infos if i.text))
    vision = _use_slide_vision(ai, options)

    renders: dict[int, bytes] = {}
    need_render = {s.idx for s in slides
                   if vision or any(b.overlays for b in s.blocks)}
    if need_render:
        _emit(progress, "Render slide")
        renders = pptx_render.render_with_powerpoint(path, len(slides), slide_w, slide_h) or {}
        for s in slides:
            if s.idx in need_render and s.idx not in renders:
                renders[s.idx] = pptx_render.render_with_pil(s.infos, slide_w, slide_h)

    for s in slides:
        _emit(progress, f"Slide {s.idx}/{len(slides)}")
        _analyze_pictures(s, ai_images, renders.get(s.idx), slide_w, slide_h, language, progress)
        if s.shape_diagram and not vision:
            _raster_shape_diagram(s, ai_images, slide_w, slide_h, progress)

    vision_md: dict[int, str] = {}
    if vision:
        vision_md = _run_slide_vision(slides, ai, renders, slide_w, slide_h, language,
                                      progress, options)

    out: list[str] = [f"# {path.stem}", ""]
    for n, s in enumerate(slides):
        if n:
            out.extend(["---", ""])
        out.append(vision_md.get(s.idx) or _deterministic_markdown(s, slide_w, slide_h))
        out.append("")
        if s.notes:
            out.extend(["### Notes", s.notes, ""])
    return "\n".join(out).rstrip() + "\n"


# ---------- building the slide ----------

def _prepare(idx: int, slide, slide_w: float, slide_h: float) -> _Slide:
    infos = collect_shapes(slide)
    s = _Slide(idx=idx, slide=slide, infos=infos, title=find_title(infos, slide_h),
               notes=_speaker_notes(slide))
    s.flow = extract_flow(infos, slide_w, slide_h)
    consumed = s.flow.consumed if s.flow else set()
    overlays = find_overlays(infos)
    overlay_ids = {t.sid for ts in overlays.values() for t in ts}
    for info in infos:
        if info is s.title or info.sid in consumed or info.sid in overlay_ids:
            continue
        if is_page_number(info, slide_h):
            continue
        box = (info.left, info.top, info.right, info.bottom)
        if info.kind == "text":
            md = frame_markdown(info.shape, default_bullets=info.role == "body")
            if md:
                s.blocks.append(Block("text", *box, markdown=md, info=info))
        elif info.kind == "table":
            md = table_markdown(info.shape)
            if md:
                s.blocks.append(Block("table", *box, markdown=md, info=info))
        elif info.kind == "chart":
            s.blocks.append(Block("chart", *box, markdown=_chart_markdown(info.shape), info=info))
        elif info.kind == "picture":
            s.blocks.append(Block("picture", *box, info=info, overlays=overlays.get(info.sid, [])))
    if s.flow:
        s.blocks.append(Block("diagram", *s.flow.bbox,
                              markdown="```mermaid\n" + s.flow.mermaid() + "\n```"))
    elif _looks_like_shape_diagram(infos) and not any(b.kind in ("picture", "chart") for b in s.blocks):
        s.shape_diagram = True
    return s


def _analyze_pictures(s: _Slide, ai: AIClient, render: bytes | None, slide_w: float,
                      slide_h: float, language: str, progress: ProgressFn | None) -> None:
    pictures = [b for b in s.blocks if b.kind == "picture"]
    context = _summarize_context(s.title_text, [b.markdown for b in s.blocks if b.kind == "text"])
    for n, b in enumerate(pictures, start=1):
        pic = b.info.shape
        if _marked_decorative(pic) or _is_office_icon(pic):
            continue
        _emit(progress, f"Slide {s.idx}: gambar {n}/{len(pictures)}")
        alt = _picture_alt_text(pic)
        if len(alt) >= _ALT_TEXT_MIN_LEN and not _GENERATED_ALT_RE.search(alt):
            b.markdown = f"> **Gambar (alt text penulis):** {' '.join(alt.split())}"
            continue
        try:
            if b.overlays and render:
                # Text boxes sit on top of this picture (e.g. a word patched
                # over the bitmap): read it the way the slide shows it.
                blob = pptx_render.crop_box(render, b.info, slide_w, slide_h)
                content_type = "image/png"
            else:
                blob = pic.image.blob
                content_type = pic.image.content_type or "image/png"
        except Exception as e:
            b.markdown = f"> *[Gagal mengambil gambar {n}: {e}]*"
            continue
        try:
            analysis = ai.analyze_image(
                blob, media_type=content_type,
                extra_context=(f"Gambar ini dari slide {s.idx} sebuah PPTX. "
                               f"Output language: {language}.\n{context}").strip(),
            )
        except Exception as e:
            b.markdown = f"> *[Gagal menganalisis gambar {n}: {e}]*"
            continue
        b.markdown = analysis.to_markdown_block(alt_hint=f"slide{s.idx}_img{n}")
    s.blocks = [b for b in s.blocks if b.kind != "picture" or b.markdown]


def _raster_shape_diagram(s: _Slide, ai: AIClient, slide_w: float, slide_h: float,
                          progress: ProgressFn | None) -> None:
    """Shape diagram whose connectors did not resolve to boxes: render the
    shapes and let the vision model describe it (pre-2026-09-30 behaviour)."""
    _emit(progress, f"Slide {s.idx}: render shape diagram")
    try:
        png = render_slide_shapes(s.slide, slide_w, slide_h)
        analysis = ai.analyze_image(
            png, media_type="image/png",
            extra_context=(f"Slide {s.idx} sebuah PPTX. Diagram digambar pakai shape + connector "
                           f"PPTX asli (bukan bitmap). Render ini rasterisasi shape-nya — hanya "
                           f"bentuk & label yang penting, abaikan warna/font.\n"
                           + _summarize_context(s.title_text, [])).strip(),
        )
    except Exception as e:
        md = f"> *[Gagal render/analisa shape-diagram slide {s.idx}: {e}]*"
    else:
        md = analysis.to_markdown_block(alt_hint=f"slide{s.idx}_shapes")
    s.blocks = [b for b in s.blocks if b.kind != "text"]
    if md:
        s.blocks.append(Block("diagram", 0, 0, slide_w, slide_h, markdown=md))


def _deterministic_markdown(s: _Slide, slide_w: float, slide_h: float) -> str:
    parts = [f"## {s.title_text}" if s.title_text else f"## Slide {s.idx}"]
    parts += [b.markdown for b in reading_order(s.blocks, slide_w, slide_h) if b.markdown]
    return "\n\n".join(parts)


# ---------- slide vision ----------

def _use_slide_vision(ai: AIClient, options: ConvertOptions | None) -> bool:
    if not hasattr(ai, "transcribe_page"):
        return False
    mode = options.pptx_slide_vision if options is not None else "auto"
    if mode == "never":
        return False
    if mode == "always":
        return True
    if options is not None and options.image_mode == "none":
        return False
    return not getattr(ai, "is_local", True)


def slide_text_layer(s: _Slide, slide_w: float, slide_h: float) -> tuple[str, dict[str, str]]:
    """(text layer for page vision, {placeholder token: exact Markdown})."""
    entries: list[str] = []
    tokens: dict[str, str] = {}
    counts: dict[str, int] = {}
    if s.title:
        entries.append(format_text_layer_entry(s.title.left / slide_w, s.title.top / slide_h,
                                               s.title.text))
    for b in reading_order(s.blocks, slide_w, slide_h):
        if not b.markdown:
            continue
        x, y = b.left / slide_w, b.top / slide_h
        if b.kind == "text":
            entries.append(format_text_layer_entry(x, y, b.markdown))
            continue
        name = _TOKEN_KINDS[b.kind]
        counts[name] = counts.get(name, 0) + 1
        token = f"[[{name} {counts[name]}]]"
        tokens[token] = b.markdown
        entries.append(format_text_layer_entry(x, y, token))
    layer = "\n".join(entries)
    if not s.flow:
        edges = resolve_edges(s.infos, slide_w, slide_h)
        if edges:
            layer += "\n\n" + "\n".join([_CONNECTORS_HEADER] + Flow(nodes=[], edges=edges).hint_lines())
    if tokens:
        layer += "\n\nPLACEHOLDERS (filled in afterwards from the file — do not reproduce):\n"
        layer += "\n".join(f"{t}: {_describe_token(md)}" for t, md in tokens.items())
    return layer, tokens


def _describe_token(md: str) -> str:
    if md.startswith("|"):
        rows = sum(1 for ln in md.splitlines() if ln.startswith("|")) - 1
        return f"table, {rows} rows"
    if md.startswith("```mermaid"):
        return "flowchart (Mermaid) of the connected boxes"
    if md.startswith("> **Gambar"):
        return "picture: " + md.split(":**", 1)[-1].strip()[:120]
    return "text transcribed from a picture"


def _run_slide_vision(slides: list[_Slide], ai: AIClient, renders: dict[int, bytes],
                      slide_w: float, slide_h: float, language: str,
                      progress: ProgressFn | None, options: ConvertOptions | None) -> dict[int, str]:
    jobs = {}
    for s in slides:
        if s.idx not in renders or not (s.blocks or s.title):
            continue
        layer, tokens = slide_text_layer(s, slide_w, slide_h)
        jobs[s.idx] = (s, layer, tokens)

    def _one(idx: int):
        s, layer, _tokens = jobs[idx]
        ctx = (f"Slide {idx} of {len(slides)}. Document language: {language}. "
               f"Slide title (from the file): {s.title_text or '(none)'}.")
        return ai.transcribe_page(renders[idx], layer, ctx)

    stats = options._stats if options is not None else None
    results: dict[int, str] = {}
    _emit(progress, f"AI per slide: {len(jobs)} dari {len(slides)} slide")
    with ThreadPoolExecutor(max_workers=_SLIDE_VISION_WORKERS) as pool:
        futures = {idx: pool.submit(_one, idx) for idx in sorted(jobs)}
        for idx, fut in futures.items():
            s, _layer, tokens = jobs[idx]
            try:
                tr = fut.result()
            except Exception as e:
                _emit(progress, f"Slide {idx}: AI per slide gagal ({e}) — susunan dari XML")
                if stats is not None:
                    stats["page_vision_fallbacks"] = stats.get("page_vision_fallbacks", 0) + 1
                continue
            if tr.coverage < PAGE_MIN_COVERAGE:
                # Words still missing after the retry: the XML layout is
                # complete by construction, so it wins over a lossy answer.
                _emit(progress, f"Slide {idx}: AI per slide kehilangan teks "
                                f"(cakupan {tr.coverage:.0%}) — susunan dari XML")
                if stats is not None:
                    stats["page_vision_fallbacks"] = stats.get("page_vision_fallbacks", 0) + 1
                continue
            _emit(progress, f"Slide {idx}/{len(slides)}: AI per slide "
                            f"(cakupan teks {tr.coverage:.0%}{', diulang' if tr.retried else ''})")
            results[idx] = finish_slide_markdown(tr.markdown, tokens, s.title_text or f"Slide {idx}")
            if stats is not None:
                stats["page_vision_pages"] = stats.get("page_vision_pages", 0) + 1
                stats["page_vision_retries"] = stats.get("page_vision_retries", 0) + int(tr.retried)
                stats["page_vision_min_coverage"] = min(stats.get("page_vision_min_coverage", 1.0),
                                                        tr.coverage)
                stats["page_vision_tokens_in"] = stats.get("page_vision_tokens_in", 0) + tr.input_tokens
                stats["page_vision_tokens_out"] = (stats.get("page_vision_tokens_out", 0)
                                                   + tr.output_tokens)
    return results


def finish_slide_markdown(md: str, tokens: dict[str, str], title: str) -> str:
    """Put the exact content in place of the placeholders, drop what the
    model reproduced on its own anyway, and make sure the slide has its `##`."""
    md = _drop_model_copies(md.strip(), tokens)
    used: set[str] = set()

    def _fill(m: re.Match[str]) -> str:
        token = f"[[{m.group(1)} {m.group(2)}]]"
        if token not in tokens or token in used:
            return ""
        used.add(token)
        return tokens[token]

    md = _TOKEN_LINE_RE.sub(_fill, md)
    md = _TOKEN_RE.sub(lambda m: "\n\n" + _fill(m) + "\n\n", md)
    missing = [content for token, content in tokens.items() if token not in used]
    if missing:
        md = md.rstrip() + "\n\n" + "\n\n".join(missing)
    md = re.sub(r"\n{3,}", "\n\n", md).strip()
    if not re.search(r"^## ", md, re.M):
        md = f"## {title}\n\n{md}"
    return md


def _drop_model_copies(md: str, tokens: dict[str, str]) -> str:
    """A model that also transcribed a placeholder's table or redrew its
    flowchart duplicates content; the file's version wins."""
    rows = {_norm_row(ln) for content in tokens.values() for ln in content.splitlines()
            if ln.startswith("|") and not re.match(r"^\|[\s:|-]+\|$", ln)}
    lines = md.splitlines()
    out: list[str] = []
    i = 0
    while i < len(lines):
        if lines[i].startswith("|") and rows:
            j = i
            while j < len(lines) and lines[j].startswith("|"):
                j += 1
            block = [ln for ln in lines[i:j] if not re.match(r"^\|[\s:|-]+\|$", ln)]
            if block and sum(_norm_row(ln) in rows for ln in block) * 2 >= len(block):
                i = j
                continue
            out.extend(lines[i:j])
            i = j
            continue
        out.append(lines[i])
        i += 1
    md = "\n".join(out)
    if any(t.startswith("[[DIAGRAM") for t in tokens):
        md = re.sub(r"```mermaid[ \t]*\n.*?```", "", md, flags=re.S)
    return md


def _norm_row(line: str) -> str:
    return "|".join(c.strip().lower() for c in line.strip().strip("|").split("|"))


# ---------- pieces kept from the earlier converter ----------

def _chart_markdown(shape) -> str:
    """Chart data via python-pptx (structured) → markdown table. Skips the AI
    entirely: chart data is deterministic and models misread axis values."""
    try:
        chart = shape.chart
    except Exception:
        return "> *[Chart: tidak bisa dibaca]*"
    title = _chart_title(chart)
    heading = f"**Chart: {title}**" if title else "**Chart**"
    try:
        plot = chart.plots[0]
        categories = [str(c) for c in plot.categories]
        series_data = [(str(s.name) if getattr(s, "name", None) else "Series", list(s.values))
                       for s in plot.series]
    except Exception as e:
        return f"{heading}\n\n> *[Gagal ekstrak data chart: {e}]*"
    if not categories or not series_data:
        return f"{heading}\n\n> *[Chart kosong atau tidak terstruktur]*"
    header = ["Category"] + [name for name, _ in series_data]
    lines = ["| " + " | ".join(header) + " |", "| " + " | ".join("---" for _ in header) + " |"]
    for i, cat in enumerate(categories):
        row = [cat] + [_fmt_num(vals[i] if i < len(vals) else "") for _, vals in series_data]
        lines.append("| " + " | ".join(row) + " |")
    return f"{heading}\n\n" + "\n".join(lines)


def _chart_title(chart) -> str:
    try:
        if chart.has_title and chart.chart_title.has_text_frame:
            return chart.chart_title.text_frame.text.strip()
    except Exception:
        pass
    return ""


def _fmt_num(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        s = f"{v:.4f}".rstrip("0").rstrip(".")
        return s or "0"
    return str(v)


def _picture_alt_text(pic) -> str:
    """Author-supplied Alt Text from `<p:cNvPr descr="...">`, or ""."""
    try:
        descrs = pic._element.xpath(".//*[local-name()='cNvPr']/@descr")
        if descrs:
            return str(descrs[0]).strip()
    except Exception:
        pass
    return ""


def _marked_decorative(pic) -> bool:
    """The author ticked "Mark as decorative" (`<adec:decorative val="1"/>`)."""
    try:
        vals = pic._element.xpath(".//*[local-name()='decorative']/@val")
    except Exception:
        return False
    return any(str(v).lower() in {"1", "true"} for v in vals)


def _is_office_icon(pic) -> bool:
    """An icon from Office's icon library: an SVG picture whose alt text is
    still the icon's name. A pictogram, never content."""
    try:
        has_svg = bool(pic._element.xpath(".//*[local-name()='svgBlip']"))
    except Exception:
        return False
    return has_svg and bool(_OFFICE_ICON_ALT_RE.match(_picture_alt_text(pic)))


def _speaker_notes(slide) -> str:
    try:
        if not slide.has_notes_slide:
            return ""
        return slide.notes_slide.notes_text_frame.text.strip()
    except Exception:
        return ""


def _summarize_context(title: str, text_blocks: list[str]) -> str:
    """Compact context for AI image calls: slide title + some body text."""
    parts: list[str] = []
    if title:
        parts.append(f"Judul slide: {title}")
    joined = " | ".join(" ".join(b.split()) for b in text_blocks if b)
    if joined:
        if len(joined) > _CONTEXT_TEXT_MAX:
            joined = joined[:_CONTEXT_TEXT_MAX] + "…"
        parts.append(f"Teks slide: {joined}")
    return "\n".join(parts)


def _looks_like_shape_diagram(infos: list[ShapeInfo]) -> bool:
    shapes = sum(1 for i in infos if i.kind in ("text", "shape") and i.prst)
    connectors = sum(1 for i in infos if i.kind == "connector")
    return shapes >= 3 and connectors >= 1


def _emit(progress: ProgressFn | None, msg: str) -> None:
    if progress:
        progress(msg)
